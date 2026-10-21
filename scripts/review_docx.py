"""Shared python-docx helpers for the Review 2 and Review 3 assessment documents."""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

NAVY = RGBColor(0x0D, 0x1B, 0x2A)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
FONT = "Calibri"

STUDENT = "Upayan Mazumder"
REGISTER_NUMBER = "24BDS0367"
PROJECT_TITLE = "AuxFlow: Static Analysis for Multi-Pass Document Compilation"
PROGRAMME = "B.Tech CSE (Data Science)"
GUIDE = "Banupriya R"


def new_document() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    return doc


def shade_cell(cell, color_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(
        qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): color_hex}
    )
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 10) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = FONT
    run.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(14)
    srun = sub.add_run(subtitle)
    srun.italic = True
    srun.font.size = Pt(12)
    srun.font.color.rgb = GREY


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(16)
    run.font.color.rgb = NAVY
    p_pr = p._p.get_or_add_pPr()
    border = p_pr.makeelement(qn("w:pBdr"), {})
    bottom = border.makeelement(
        qn("w:bottom"),
        {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "4", qn("w:color"): "1B3A57"},
    )
    border.append(bottom)
    p_pr.append(border)


def add_para(doc: Document, text: str, bold: bool = False) -> Paragraph:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(11)
    run.bold = bold
    return p


def add_label(doc: Document, text: str) -> Paragraph:
    """A bold subheading kept on the same page as the content after it."""
    p = add_para(doc, text, bold=True)
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.bold = True
        lead.font.name = FONT
        lead.font.size = Pt(11)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(11)


def add_details_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(10.5)
    for row, (label, value) in zip(table.rows, rows, strict=True):
        set_cell_text(row.cells[0], label, bold=True, size=11)
        shade_cell(row.cells[0], "E6EDF3")
        set_cell_text(row.cells[1], value, size=11)
    doc.add_paragraph()


def add_grid(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows) + 1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        set_cell_text(cell, text, bold=True, size=9)
        shade_cell(cell, "E6EDF3")
    for row, values in zip(table.rows[1:], rows, strict=True):
        for cell, text in zip(row.cells, values, strict=True):
            set_cell_text(cell, text, size=9)
    doc.add_paragraph()


def add_figure(doc: Document, path: str, caption: str, width_cm: float = 15.5) -> None:
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = GREY
